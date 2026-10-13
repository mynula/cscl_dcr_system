import csv
import io
from typing import List, Optional
from zipfile import BadZipFile
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile, status
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import EmailStr
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.user import User
from app.models.master import Territory, Hospital, Doctor, VisitType
from app.schemas.master import (
    TerritoryCreate, TerritoryResponse,
    HospitalCreate, HospitalResponse,
    DoctorCreate, DoctorResponse,
    VisitTypeCreate, VisitTypeResponse
)
from app.schemas.user import UserCreate, UserResponse
from app.services.auth_service import (
    get_current_admin_only_user,
    get_current_admin_user,
    get_password_hash,
)

router = APIRouter(prefix="/admin", tags=["Master Admin CRUD"])
templates = Jinja2Templates(directory="app/templates")
MAX_MASTER_IMPORT_BYTES = 5 * 1024 * 1024
MAX_MASTER_IMPORT_ROWS = 5000
MASTER_IMPORT_HEADERS = {
    "territories": ["code", "name", "region", "description", "is_active"],
    "hospitals": ["code", "name", "type", "territory_code", "address", "phone", "is_active"],
    "doctors": [
        "name", "specialty", "territory_code", "hospital_code", "qualification",
        "designation", "phone", "email", "chamber_address", "is_active",
    ],
    "visit-types": ["name", "description", "is_active"],
}
MASTER_CONFIG = {
    "territories": (Territory, "territories"),
    "hospitals": (Hospital, "hospitals"),
    "doctors": (Doctor, "doctors"),
    "visit-types": (VisitType, "visit_types"),
}


def _master_model(resource: str):
    try:
        return MASTER_CONFIG[resource]
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Unknown master-data type.") from exc


def _get_master_record(db: Session, resource: str, record_id: int):
    model, _ = _master_model(resource)
    record = db.query(model).filter(model.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="Master-data record not found.")
    return record


def _render_master_table(request: Request, db: Session, resource: str):
    model, tab = _master_model(resource)
    rows = db.query(model).order_by(model.name if model is not Territory else model.code).all()
    name = {
        "territories": "territory_table.html",
        "hospitals": "hospital_table.html",
        "doctors": "doctor_table.html",
        "visit-types": "visit_type_table.html",
    }[resource]
    return templates.TemplateResponse(
        request=request,
        name=f"admin/partials/{name}",
        context={tab: rows},
    )


def _clean_form_value(form, field: str, optional: bool = False):
    value = str(form.get(field) or "").strip()
    return value or (None if optional else "")


def _parse_active(value: object, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    normalized = str(value).strip().casefold()
    if normalized in {"true", "1", "yes", "active", "y"}:
        return True
    if normalized in {"false", "0", "no", "inactive", "n"}:
        return False
    raise ValueError("is_active must be true/false, yes/no, 1/0, or active/inactive.")


def _read_master_upload(filename: str, content: bytes, resource: str):
    if len(content) > MAX_MASTER_IMPORT_BYTES:
        raise ValueError("Import file exceeds the 5 MB limit.")
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension not in {"csv", "xlsx"}:
        raise ValueError("Upload a .csv or .xlsx file.")

    workbook = None
    try:
        if extension == "csv":
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            headers = [str(value or "").strip().casefold() for value in (reader.fieldnames or [])]
            rows = [
                {
                    headers[index]: value
                    for index, value in enumerate(row.values())
                    if index < len(headers)
                }
                for row in reader
            ]
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet.max_row > MAX_MASTER_IMPORT_ROWS + 1:
                raise ValueError(f"Upload no more than {MAX_MASTER_IMPORT_ROWS} rows.")
            values = sheet.iter_rows(values_only=True)
            header_row = next(values, ())
            headers = [str(value or "").strip().casefold() for value in header_row]
            rows = [
                {
                    headers[index]: value
                    for index, value in enumerate(row)
                    if index < len(headers) and headers[index]
                }
                for row in values
                if any(value is not None for value in row)
            ]
    except ValueError:
        raise
    except (UnicodeDecodeError, OSError, BadZipFile, InvalidFileException, EOFError) as exc:
        raise ValueError("Could not read the uploaded CSV/Excel file.") from exc
    finally:
        if workbook:
            workbook.close()

    missing = [header for header in MASTER_IMPORT_HEADERS[resource] if header not in headers]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))
    if not rows:
        raise ValueError("The uploaded file contains no data rows.")
    if len(rows) > MAX_MASTER_IMPORT_ROWS:
        raise ValueError(f"Upload no more than {MAX_MASTER_IMPORT_ROWS} rows.")
    return rows


def _validate_master_import_row(db: Session, resource: str, values: dict, seen: set):
    def required(key: str) -> str:
        value = str(values.get(key) or "").strip()
        if not value:
            raise ValueError(f"{key} is required.")
        return value

    active = _parse_active(values.get("is_active"))
    if resource == "territories":
        code = required("code").upper()
        name = required("name")
        region = required("region")
        if code.casefold() in seen or db.query(Territory.id).filter(func.upper(Territory.code) == code).first():
            raise ValueError(f"Territory code '{code}' already exists.")
        seen.add(code.casefold())
        return Territory(
            code=code, name=name, region=region,
            description=str(values.get("description") or "").strip() or None,
            is_active=active,
        )
    if resource == "hospitals":
        code = required("code").upper()
        name = required("name")
        territory_code = required("territory_code").upper()
        territory = db.query(Territory).filter(
            func.upper(Territory.code) == territory_code,
            Territory.is_active.is_(True),
        ).first()
        if not territory:
            raise ValueError(f"Active territory code '{territory_code}' was not found.")
        if code.casefold() in seen or db.query(Hospital.id).filter(func.upper(Hospital.code) == code).first():
            raise ValueError(f"Hospital code '{code}' already exists.")
        seen.add(code.casefold())
        return Hospital(
            code=code,
            name=name,
            type=str(values.get("type") or "General Hospital").strip(),
            territory_id=territory.id,
            address=str(values.get("address") or "").strip() or None,
            phone=str(values.get("phone") or "").strip() or None,
            is_active=active,
        )
    if resource == "doctors":
        name = required("name")
        specialty = required("specialty")
        territory_code = required("territory_code").upper()
        territory = db.query(Territory).filter(
            func.upper(Territory.code) == territory_code,
            Territory.is_active.is_(True),
        ).first()
        if not territory:
            raise ValueError(f"Active territory code '{territory_code}' was not found.")
        hospital_code = str(values.get("hospital_code") or "").strip().upper()
        hospital_id = None
        if hospital_code:
            hospital = db.query(Hospital).filter(
                func.upper(Hospital.code) == hospital_code,
                Hospital.territory_id == territory.id,
                Hospital.is_active.is_(True),
            ).first()
            if not hospital:
                raise ValueError(
                    f"Active hospital code '{hospital_code}' was not found in territory '{territory_code}'."
                )
            hospital_id = hospital.id
        signature = (name.casefold(), territory.id, hospital_id)
        if signature in seen or db.query(Doctor.id).filter(
            func.lower(Doctor.name) == name.casefold(),
            Doctor.territory_id == territory.id,
            Doctor.hospital_id == hospital_id,
        ).first():
            raise ValueError(f"Doctor '{name}' is already in that territory/hospital.")
        seen.add(signature)
        return Doctor(
            name=name,
            specialty=specialty,
            territory_id=territory.id,
            hospital_id=hospital_id,
            qualification=str(values.get("qualification") or "").strip() or None,
            designation=str(values.get("designation") or "").strip() or None,
            phone=str(values.get("phone") or "").strip() or None,
            email=str(values.get("email") or "").strip() or None,
            chamber_address=str(values.get("chamber_address") or "").strip() or None,
            is_active=active,
        )
    if resource == "visit-types":
        name = required("name")
        signature = name.casefold()
        if signature in seen or db.query(VisitType.id).filter(
            func.lower(VisitType.name) == signature
        ).first():
            raise ValueError(f"Visit type '{name}' already exists.")
        seen.add(signature)
        return VisitType(
            name=name,
            description=str(values.get("description") or "").strip() or None,
            is_active=active,
        )
    raise ValueError("Unknown master-data type.")


def _validate_master_field_lengths(record):
    for column in record.__table__.columns:
        value = getattr(record, column.name, None)
        if isinstance(value, str) and column.type.length and len(value) > column.type.length:
            raise ValueError(
                f"{column.name} must be no longer than {column.type.length} characters."
            )


# --- Admin Dashboard View ---
@router.get("", response_class=HTMLResponse)
@router.get("/", response_class=HTMLResponse)
async def admin_dashboard(
    request: Request,
    tab: str = "territories",
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    territories = db.query(Territory).order_by(Territory.name).all()
    hospitals = db.query(Hospital).order_by(Hospital.name).all()
    doctors = db.query(Doctor).order_by(Doctor.name).all()
    visit_types = db.query(VisitType).order_by(VisitType.name).all()
    users = (
        db.query(User).order_by(User.full_name).all()
        if admin_user.role == "admin"
        else []
    )

    return templates.TemplateResponse(
        request=request,
        name="admin/index.html",
        context={
            "user": admin_user,
            "territories": territories,
            "hospitals": hospitals,
            "doctors": doctors,
            "visit_types": visit_types,
            "users": users,
            "active_tab": tab if tab in {"territories", "hospitals", "doctors", "visit_types", "users"} else "territories"
        }
    )


@router.get("/import/{resource}/template.csv")
async def download_master_import_template(
    resource: str,
    admin_user: User = Depends(get_current_admin_user),
):
    headers = MASTER_IMPORT_HEADERS.get(resource)
    if not headers:
        raise HTTPException(status_code=404, detail="Unknown master-data type.")
    examples = {
        "territories": ["TERR-DHK-05", "Mirpur North", "Dhaka North", "Coverage area description", "true"],
        "hospitals": ["HOSP-DMCH-01", "Dhaka Medical College Hospital", "Government Medical College", "TERR-DHK-01", "Dhaka", "+880", "true"],
        "doctors": ["Dr Example", "Cardiology", "TERR-DHK-01", "HOSP-DMCH-01", "MBBS, FCPS", "Consultant", "01XXXXXXXXX", "doctor@example.com", "Dhaka", "true"],
        "visit-types": ["Clinical Seminar", "Discussion or instruction", "true"],
    }
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerow(examples[resource])
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{resource}-template.csv"'},
    )


@router.get("/import/{resource}", response_class=HTMLResponse)
async def master_import_page(
    request: Request,
    resource: str,
    admin_user: User = Depends(get_current_admin_user),
):
    _master_model(resource)
    return templates.TemplateResponse(
        request=request,
        name="admin/master_import.html",
        context={
            "user": admin_user,
            "resource": resource,
            "resource_label": MASTER_CONFIG[resource][1].replace("_", " ").title(),
            "headers": MASTER_IMPORT_HEADERS[resource],
        },
    )


@router.post("/import/{resource}", response_class=HTMLResponse)
async def import_master_data(
    request: Request,
    resource: str,
    upload: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    _master_model(resource)
    try:
        content = await upload.read(MAX_MASTER_IMPORT_BYTES + 1)
        rows = _read_master_upload(upload.filename or "", content, resource)
        seen = set()
        records = []
        errors = []
        for row_number, values in enumerate(rows, start=2):
            try:
                record = _validate_master_import_row(db, resource, values, seen)
                _validate_master_field_lengths(record)
                records.append(record)
            except (ValueError, TypeError) as exc:
                errors.append({"row_number": row_number, "message": str(exc)})
        if errors:
            return templates.TemplateResponse(
                request=request,
                name="admin/master_import_result.html",
                context={
                    "user": admin_user,
                    "resource": resource,
                    "resource_label": MASTER_CONFIG[resource][1].replace("_", " ").title(),
                    "errors": errors,
                    "created_count": 0,
                },
                status_code=400,
            )
        db.add_all(records)
        db.commit()
    except ValueError as exc:
        db.rollback()
        return templates.TemplateResponse(
            request=request,
            name="admin/master_import_result.html",
            context={
                "user": admin_user,
                "resource": resource,
                "resource_label": MASTER_CONFIG[resource][1].replace("_", " ").title(),
                "errors": [{"row_number": None, "message": str(exc)}],
                "created_count": 0,
            },
            status_code=400,
        )
    except IntegrityError as exc:
        db.rollback()
        return templates.TemplateResponse(
            request=request,
            name="admin/master_import_result.html",
            context={
                "user": admin_user,
                "resource": resource,
                "resource_label": MASTER_CONFIG[resource][1].replace("_", " ").title(),
                "errors": [{"row_number": None, "message": "A duplicate or invalid reference prevented the import. No rows were saved."}],
                "created_count": 0,
            },
            status_code=400,
        )
    return templates.TemplateResponse(
        request=request,
        name="admin/master_import_result.html",
        context={
            "user": admin_user,
            "resource": resource,
            "resource_label": MASTER_CONFIG[resource][1].replace("_", " ").title(),
            "errors": [],
            "created_count": len(records),
        },
    )


@router.get("/{resource}/{record_id}/edit", response_class=HTMLResponse)
async def edit_master_record_page(
    request: Request,
    resource: str,
    record_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    record = _get_master_record(db, resource, record_id)
    return templates.TemplateResponse(
        request=request,
        name="admin/master_edit.html",
        context={
            "user": admin_user,
            "resource": resource,
            "tab": MASTER_CONFIG[resource][1],
            "record": record,
            "territories": db.query(Territory).order_by(Territory.name).all(),
            "hospitals": db.query(Hospital).order_by(Hospital.name).all(),
        },
    )


@router.post("/{resource}/{record_id}/edit")
async def update_master_record(
    request: Request,
    resource: str,
    record_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    record = _get_master_record(db, resource, record_id)
    form = await request.form()
    try:
        if resource == "territories":
            code = _clean_form_value(form, "code").upper()
            name = _clean_form_value(form, "name")
            region = _clean_form_value(form, "region")
            if not code or not name or not region:
                raise ValueError("Code, name, and region are required.")
            conflict = db.query(Territory.id).filter(
                func.upper(Territory.code) == code, Territory.id != record_id
            ).first()
            if conflict:
                raise ValueError(f"Territory code '{code}' already exists.")
            record.code, record.name, record.region = code, name, region
            record.description = _clean_form_value(form, "description", optional=True)
        elif resource == "hospitals":
            code = _clean_form_value(form, "code").upper()
            name = _clean_form_value(form, "name")
            try:
                territory_id = int(_clean_form_value(form, "territory_id"))
            except ValueError as exc:
                raise ValueError("Select an existing territory.") from exc
            territory = db.query(Territory).filter(Territory.id == territory_id).first()
            if not territory:
                raise ValueError("Select an existing territory.")
            if not code or not name:
                raise ValueError("Code and name are required.")
            conflict = db.query(Hospital.id).filter(
                func.upper(Hospital.code) == code, Hospital.id != record_id
            ).first()
            if conflict:
                raise ValueError(f"Hospital code '{code}' already exists.")
            record.code, record.name = code, name
            record.type = _clean_form_value(form, "type") or "General Hospital"
            record.territory_id = territory_id
            record.address = _clean_form_value(form, "address", optional=True)
            record.phone = _clean_form_value(form, "phone", optional=True)
        elif resource == "doctors":
            name = _clean_form_value(form, "name")
            specialty = _clean_form_value(form, "specialty")
            try:
                territory_id = int(_clean_form_value(form, "territory_id"))
            except ValueError as exc:
                raise ValueError("Select an existing territory.") from exc
            if not name or not specialty:
                raise ValueError("Name and specialty are required.")
            if not db.query(Territory.id).filter(Territory.id == territory_id).first():
                raise ValueError("Select an existing territory.")
            hospital_raw = _clean_form_value(form, "hospital_id", optional=True)
            hospital_id = int(hospital_raw) if hospital_raw else None
            if hospital_id and not db.query(Hospital.id).filter(
                Hospital.id == hospital_id, Hospital.territory_id == territory_id
            ).first():
                raise ValueError("The selected hospital must belong to the selected territory.")
            record.name, record.specialty = name, specialty
            record.territory_id, record.hospital_id = territory_id, hospital_id
            for field in ("qualification", "designation", "phone", "email", "chamber_address"):
                setattr(record, field, _clean_form_value(form, field, optional=True))
        elif resource == "visit-types":
            name = _clean_form_value(form, "name")
            if not name:
                raise ValueError("Name is required.")
            conflict = db.query(VisitType.id).filter(
                func.lower(VisitType.name) == name.casefold(),
                VisitType.id != record_id,
            ).first()
            if conflict:
                raise ValueError(f"Visit type '{name}' already exists.")
            record.name = name
            record.description = _clean_form_value(form, "description", optional=True)
        else:
            raise HTTPException(status_code=404, detail="Unknown master-data type.")

        record.is_active = _parse_active(form.get("is_active"))
        _validate_master_field_lengths(record)
        db.commit()
    except ValueError as exc:
        db.rollback()
        record = _get_master_record(db, resource, record_id)
        return templates.TemplateResponse(
            request=request,
            name="admin/master_edit.html",
            context={
                "user": admin_user,
                "resource": resource,
                "tab": MASTER_CONFIG[resource][1],
                "record": record,
                "territories": db.query(Territory).order_by(Territory.name).all(),
                "hospitals": db.query(Hospital).order_by(Hospital.name).all(),
                "error": str(exc),
            },
            status_code=400,
        )
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail="The update conflicts with an existing master-data record.") from exc
    return RedirectResponse(
        url=f"/admin?tab={MASTER_CONFIG[resource][1]}",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.post("/{resource}/{record_id}/active")
async def update_master_record_status(
    request: Request,
    resource: str,
    record_id: int,
    is_active: str = Form(...),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user),
):
    record = _get_master_record(db, resource, record_id)
    try:
        record.is_active = _parse_active(is_active, default=False)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.commit()
    if request.headers.get("hx-request"):
        return _render_master_table(request, db, resource)
    return RedirectResponse(
        url=f"/admin?tab={MASTER_CONFIG[resource][1]}",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.get("/users", response_model=List[UserResponse])
async def list_users(
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_only_user)
):
    """List user accounts for administrator account management."""
    return db.query(User).order_by(User.full_name).all()


@router.post("/users")
async def create_user(
    request: Request,
    username: str = Form(..., max_length=50),
    email: EmailStr = Form(..., max_length=100),
    full_name: str = Form(..., max_length=100),
    password: str = Form(..., min_length=8),
    role: str = Form(..., max_length=20),
    employee_id: Optional[str] = Form(None, max_length=30),
    phone: Optional[str] = Form(None, max_length=20),
    territory_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_only_user)
):
    """Create a user account; only administrators may assign roles."""
    username = username.strip()
    email = str(email).strip().lower()
    full_name = full_name.strip()
    role = role.strip().lower()
    employee_id = employee_id.strip() if employee_id else None
    phone = phone.strip() if phone else None

    if not username or not email or not full_name:
        raise HTTPException(status_code=400, detail="Username, email, and full name are required.")
    if role not in {"admin", "manager", "marketer"}:
        raise HTTPException(status_code=400, detail="Role must be admin, manager, or marketer.")
    if db.query(User).filter(User.username == username).first():
        raise HTTPException(status_code=400, detail=f"Username '{username}' already exists.")
    if db.query(User).filter(User.email == email).first():
        raise HTTPException(status_code=400, detail=f"Email '{email}' already exists.")
    if employee_id and db.query(User).filter(User.employee_id == employee_id).first():
        raise HTTPException(status_code=400, detail=f"Employee ID '{employee_id}' already exists.")
    if territory_id and not db.query(Territory).filter(
        Territory.id == territory_id, Territory.is_active.is_(True)
    ).first():
        raise HTTPException(status_code=400, detail="Selected territory is not active.")

    user = User(
        username=username,
        email=email,
        full_name=full_name,
        hashed_password=get_password_hash(password),
        role=role,
        employee_id=employee_id,
        phone=phone,
        territory_id=territory_id,
        is_active=True,
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    if request.headers.get("hx-request"):
        users = db.query(User).order_by(User.full_name).all()
        return templates.TemplateResponse(
            request=request,
            name="admin/partials/user_table.html",
            context={"users": users}
        )
    return UserResponse.model_validate(user)


# ==============================================================================
# TERRITORY CRUD
# ==============================================================================
@router.get("/territories", response_model=List[TerritoryResponse])
async def list_territories(
    request: Request,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    territories = db.query(Territory).order_by(Territory.name).all()
    if "text/html" in request.headers.get("accept", "") and not request.headers.get("hx-request"):
        return RedirectResponse(url="/admin", status_code=status.HTTP_302_FOUND)
    return territories


@router.post("/territories")
async def create_territory(
    request: Request,
    code: str = Form(...),
    name: str = Form(...),
    region: str = Form(...),
    description: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    code = code.strip().upper()
    name = name.strip()
    region = region.strip()
    if not code or not name or not region:
        raise HTTPException(status_code=400, detail="Code, name, and region are required.")
    existing = db.query(Territory).filter(func.upper(Territory.code) == code).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Territory code '{code}' already exists.")

    territory = Territory(
        code=code,
        name=name,
        region=region,
        description=description.strip() if description else None
    )
    db.add(territory)
    db.commit()
    db.refresh(territory)

    if request.headers.get("hx-request"):
        territories = db.query(Territory).order_by(Territory.name).all()
        return templates.TemplateResponse(
            request=request,
            name="admin/partials/territory_table.html",
            context={"territories": territories}
        )
    return RedirectResponse(url="/admin", status_code=status.HTTP_303_SEE_OTHER)


@router.delete("/territories/{territory_id}")
async def delete_territory(
    request: Request,
    territory_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    territory = db.query(Territory).filter(Territory.id == territory_id).first()
    if not territory:
        raise HTTPException(status_code=404, detail="Territory not found")
    db.delete(territory)
    db.commit()
    if request.headers.get("hx-request"):
        return _render_master_table(request, db, "territories")
    return Response(content="", status_code=status.HTTP_200_OK)


# ==============================================================================
# HOSPITAL CRUD
# ==============================================================================
@router.get("/hospitals", response_model=List[HospitalResponse])
async def list_hospitals(
    territory_id: Optional[int] = None,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    query = db.query(Hospital)
    if territory_id:
        query = query.filter(Hospital.territory_id == territory_id)
    return query.order_by(Hospital.name).all()


@router.post("/hospitals")
async def create_hospital(
    request: Request,
    code: str = Form(...),
    name: str = Form(...),
    type: str = Form("General Hospital"),
    territory_id: int = Form(...),
    address: Optional[str] = Form(None),
    phone: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    code = code.strip().upper()
    if not code or not name.strip():
        raise HTTPException(status_code=400, detail="Hospital code and name are required.")
    if not db.query(Territory.id).filter(
        Territory.id == territory_id, Territory.is_active.is_(True)
    ).first():
        raise HTTPException(status_code=400, detail="Select an active territory.")
    existing = db.query(Hospital).filter(func.upper(Hospital.code) == code).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Hospital code '{code}' already exists.")

    hospital = Hospital(
        code=code,
        name=name.strip(),
        type=type.strip(),
        territory_id=territory_id,
        address=address.strip() if address else None,
        phone=phone.strip() if phone else None
    )
    db.add(hospital)
    db.commit()
    db.refresh(hospital)

    if request.headers.get("hx-request"):
        hospitals = db.query(Hospital).order_by(Hospital.name).all()
        return templates.TemplateResponse(
            request=request,
            name="admin/partials/hospital_table.html",
            context={"hospitals": hospitals}
        )
    return RedirectResponse(url="/admin", status_code=status.HTTP_303_SEE_OTHER)


@router.delete("/hospitals/{hospital_id}")
async def delete_hospital(
    request: Request,
    hospital_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    hospital = db.query(Hospital).filter(Hospital.id == hospital_id).first()
    if not hospital:
        raise HTTPException(status_code=404, detail="Hospital not found")
    db.delete(hospital)
    db.commit()
    if request.headers.get("hx-request"):
        return _render_master_table(request, db, "hospitals")
    return Response(content="", status_code=status.HTTP_200_OK)


# ==============================================================================
# DOCTOR CRUD
# ==============================================================================
@router.get("/doctors", response_model=List[DoctorResponse])
async def list_doctors(
    territory_id: Optional[int] = None,
    hospital_id: Optional[int] = None,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    query = db.query(Doctor)
    if territory_id:
        query = query.filter(Doctor.territory_id == territory_id)
    if hospital_id:
        query = query.filter(Doctor.hospital_id == hospital_id)
    return query.order_by(Doctor.name).all()


@router.post("/doctors")
async def create_doctor(
    request: Request,
    name: str = Form(...),
    specialty: str = Form(...),
    territory_id: int = Form(...),
    qualification: Optional[str] = Form(None),
    designation: Optional[str] = Form(None),
    hospital_id: Optional[int] = Form(None),
    phone: Optional[str] = Form(None),
    email: Optional[str] = Form(None),
    chamber_address: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    if not name.strip() or not specialty.strip():
        raise HTTPException(status_code=400, detail="Doctor name and specialty are required.")
    if not db.query(Territory.id).filter(
        Territory.id == territory_id, Territory.is_active.is_(True)
    ).first():
        raise HTTPException(status_code=400, detail="Select an active territory.")
    if hospital_id and hospital_id > 0 and not db.query(Hospital.id).filter(
        Hospital.id == hospital_id,
        Hospital.territory_id == territory_id,
        Hospital.is_active.is_(True),
    ).first():
        raise HTTPException(status_code=400, detail="Select an active hospital in the selected territory.")
    doctor = Doctor(
        name=name.strip(),
        specialty=specialty.strip(),
        territory_id=territory_id,
        qualification=qualification.strip() if qualification else None,
        designation=designation.strip() if designation else None,
        hospital_id=hospital_id if hospital_id and hospital_id > 0 else None,
        phone=phone.strip() if phone else None,
        email=email.strip() if email else None,
        chamber_address=chamber_address.strip() if chamber_address else None
    )
    db.add(doctor)
    db.commit()
    db.refresh(doctor)

    if request.headers.get("hx-request"):
        doctors = db.query(Doctor).order_by(Doctor.name).all()
        return templates.TemplateResponse(
            request=request,
            name="admin/partials/doctor_table.html",
            context={"doctors": doctors}
        )
    return RedirectResponse(url="/admin", status_code=status.HTTP_303_SEE_OTHER)


@router.delete("/doctors/{doctor_id}")
async def delete_doctor(
    request: Request,
    doctor_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    doctor = db.query(Doctor).filter(Doctor.id == doctor_id).first()
    if not doctor:
        raise HTTPException(status_code=404, detail="Doctor not found")
    db.delete(doctor)
    db.commit()
    if request.headers.get("hx-request"):
        return _render_master_table(request, db, "doctors")
    return Response(content="", status_code=status.HTTP_200_OK)


# ==============================================================================
# VISIT TYPE CRUD
# ==============================================================================
@router.get("/visit-types", response_model=List[VisitTypeResponse])
async def list_visit_types(
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    return db.query(VisitType).order_by(VisitType.name).all()


@router.post("/visit-types")
async def create_visit_type(
    request: Request,
    name: str = Form(...),
    description: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    name = name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Visit type name is required.")
    existing = db.query(VisitType).filter(func.lower(VisitType.name) == name.casefold()).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Visit type '{name}' already exists.")

    visit_type = VisitType(
        name=name.strip(),
        description=description.strip() if description else None,
        is_active=True
    )
    db.add(visit_type)
    db.commit()
    db.refresh(visit_type)

    if request.headers.get("hx-request"):
        visit_types = db.query(VisitType).order_by(VisitType.name).all()
        return templates.TemplateResponse(
            request=request,
            name="admin/partials/visit_type_table.html",
            context={"visit_types": visit_types}
        )
    return RedirectResponse(url="/admin", status_code=status.HTTP_303_SEE_OTHER)


@router.delete("/visit-types/{visit_type_id}")
async def delete_visit_type(
    request: Request,
    visit_type_id: int,
    db: Session = Depends(get_db),
    admin_user: User = Depends(get_current_admin_user)
):
    visit_type = db.query(VisitType).filter(VisitType.id == visit_type_id).first()
    if not visit_type:
        raise HTTPException(status_code=404, detail="Visit Type not found")
    db.delete(visit_type)
    db.commit()
    if request.headers.get("hx-request"):
        return _render_master_table(request, db, "visit-types")
    return Response(content="", status_code=status.HTTP_200_OK)
