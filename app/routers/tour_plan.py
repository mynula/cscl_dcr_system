import csv
import html
import io
import json
import logging
import uuid
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from typing import Optional
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload, selectinload

from app.database import get_db
from app.models.dcr import FactDCR
from app.models.master import Doctor, Territory, VisitType
from app.models.tour_plan import (
    TourPlan,
    TourPlanImport,
    TourPlanTemplate,
    TourPlanVisit,
)
from app.models.user import User
from app.services.auth_service import get_current_user
from app.services.tour_plan_service import (
    MAX_PLAN_VISITS,
    PLAN_STATUSES,
    VISIT_STATUSES,
    TourPlanValidationError,
    audit_trail,
    create_plan,
    get_plan_for_user,
    get_plan_report_rows,
    get_tour_plan_setting,
    log_plan_action,
    plan_visit_actuals,
    save_plan_template,
    update_plan,
    parse_integer,
    validate_plan_period,
    validate_plan_visits,
    visible_plans_query,
)

router = APIRouter(prefix="/tour-plan", tags=["Tour Plans"])
templates = Jinja2Templates(directory="app/templates")
logger = logging.getLogger(__name__)
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
IMPORT_TTL = timedelta(hours=2)
IMPORT_HEADERS = {
    "planned_date",
    "sequence",
    "doctor_id",
    "visit_type_id",
    "planned_products",
    "planned_indication",
    "notes",
}


def _is_admin(user: User) -> bool:
    return user.role == "admin"


def _can_self_service(db: Session) -> bool:
    return get_tour_plan_setting(db).marketer_self_service


def _require_can_create(db: Session, user: User) -> None:
    if user.role == "admin":
        return
    if user.role != "marketer" or not _can_self_service(db):
        raise HTTPException(status_code=403, detail="Tour Plan self-service is not enabled for your account.")
    if not user.territory_id:
        raise HTTPException(status_code=403, detail="An Admin must assign your territory before you can create plans.")


def _require_plan_edit(db: Session, plan: TourPlan, user: User) -> None:
    if user.role == "admin":
        return
    if user.role != "marketer" or plan.marketer_id != user.id:
        raise HTTPException(status_code=403, detail="You cannot edit another marketer's plan.")
    if not _can_self_service(db):
        raise HTTPException(status_code=403, detail="Tour Plan self-service is disabled.")
    if plan.status not in {"Draft", "Returned"}:
        raise HTTPException(status_code=403, detail="Only draft or returned plans can be edited.")


def _get_plan_or_404(db: Session, plan_id: int, user: User) -> TourPlan:
    plan = get_plan_for_user(db, plan_id, user)
    if not plan:
        raise HTTPException(status_code=404, detail="Tour Plan not found.")
    return plan


def _parse_date(value: str | None, label: str) -> date:
    try:
        return date.fromisoformat(value or "")
    except ValueError as exc:
        raise TourPlanValidationError(f"{label} must be a valid date.") from exc


def _parse_visits(value: str) -> list[dict]:
    try:
        parsed = json.loads(value)
    except (TypeError, json.JSONDecodeError) as exc:
        raise TourPlanValidationError("Planned visit data is invalid; please reload the form.") from exc
    if not isinstance(parsed, list):
        raise TourPlanValidationError("Planned visit data must be a list.")
    return parsed


def _submitted_form_values(
    title: str,
    period_type: str,
    period_start: str,
    period_end: str,
    territory_id: int,
    marketer_id: int,
    visits_json: str,
) -> dict:
    try:
        rows = _parse_visits(visits_json)
    except TourPlanValidationError:
        rows = []
    return {
        "title": title,
        "period_type": period_type,
        "period_start": period_start,
        "period_end": period_end,
        "territory_id": territory_id,
        "marketer_id": marketer_id,
        "rows": rows,
    }


def _plan_form_context(
    request: Request,
    db: Session,
    user: User,
    plan: TourPlan | None = None,
    copy_from: TourPlan | None = None,
    template_rows: list[dict] | None = None,
    error: str | None = None,
    form_values: dict | None = None,
):
    actor_marketer = user.role == "marketer"
    allowed_territories = db.query(Territory).order_by(Territory.name).all()
    if actor_marketer and user.territory_id:
        allowed_territories = [
            territory for territory in allowed_territories
            if territory.id == user.territory_id
        ]
    marketer_accounts = db.query(User).filter(
        User.role == "marketer",
        User.is_active.is_(True),
    ).order_by(User.full_name).all()
    selected_territory_id = form_values.get("territory_id") if form_values else None
    selected_territory_id = selected_territory_id or (
        plan.territory_id if plan else
        copy_from.territory_id if copy_from else
        allowed_territories[0].id if actor_marketer and allowed_territories else None
    )
    doctors = db.query(Doctor).order_by(Doctor.name).all()
    visit_types = db.query(VisitType).filter(
        VisitType.is_active.is_(True)
    ).order_by(VisitType.name).all()
    plan_rows = template_rows
    if plan:
        plan_rows = [
            {
                "planned_date": visit.planned_date.isoformat(),
                "sequence": visit.sequence,
                "doctor_id": visit.doctor_id,
                "visit_type_id": visit.visit_type_id,
                "planned_products": visit.planned_products or "",
                "planned_indication": visit.planned_indication or "",
                "notes": visit.notes or "",
            }
            for visit in plan.visits
        ]
    elif copy_from:
        if copy_from.period_type == "weekly":
            target_start = copy_from.period_start + timedelta(days=7)
        else:
            year = copy_from.period_start.year + (copy_from.period_start.month == 12)
            month = 1 if copy_from.period_start.month == 12 else copy_from.period_start.month + 1
            target_start = date(year, month, 1)
        plan_rows = [
            {
                "planned_date": (target_start + (visit.planned_date - copy_from.period_start)).isoformat(),
                "sequence": visit.sequence,
                "doctor_id": visit.doctor_id,
                "visit_type_id": visit.visit_type_id,
                "planned_products": visit.planned_products or "",
                "planned_indication": visit.planned_indication or "",
                "notes": visit.notes or "",
            }
            for visit in copy_from.visits
        ]
    if form_values:
        period_start = form_values.get("period_start")
        period_end = form_values.get("period_end")
        period_type = form_values.get("period_type")
        title = form_values.get("title")
    elif plan:
        period_start = plan.period_start
        period_end = plan.period_end
        period_type = plan.period_type
        title = plan.title
    elif copy_from:
        if copy_from.period_type == "weekly":
            period_start = copy_from.period_start + timedelta(days=7)
            period_end = period_start + timedelta(days=(copy_from.period_end - copy_from.period_start).days)
        else:
            year = copy_from.period_start.year + (copy_from.period_start.month == 12)
            month = 1 if copy_from.period_start.month == 12 else copy_from.period_start.month + 1
            period_start = date(year, month, 1)
            period_end = date(year, month, monthrange(year, month)[1])
        period_type = copy_from.period_type
        title = f"{copy_from.title} (copy)"
    else:
        period_start = date.today() - timedelta(days=date.today().weekday())
        period_end = period_start + timedelta(days=4)
        period_type = "weekly"
        title = ""
    return {
        "request": request,
        "user": user,
        "plan": plan,
        "copy_from": copy_from,
        "territories": allowed_territories,
        "marketers": marketer_accounts,
        "doctors": doctors,
        "visit_types": visit_types,
        "rows": (form_values.get("rows") if form_values else None) or plan_rows or [{}],
        "selected_marketer_id": form_values.get("marketer_id") if form_values else None,
        "selected_territory_id": selected_territory_id,
        "period_start": period_start.isoformat() if isinstance(period_start, date) else period_start,
        "period_end": period_end.isoformat() if isinstance(period_end, date) else period_end,
        "period_type": period_type,
        "title": title,
        "self_service": _can_self_service(db),
        "error": error,
    }


def _decode_import_upload(filename: str, content: bytes) -> tuple[list[dict], list[str]]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise TourPlanValidationError("Import file exceeds the 2 MB limit.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in {"csv", "xlsx"}:
        raise TourPlanValidationError("Upload a .csv or .xlsx file.")
    workbook = None
    try:
        if ext == "csv":
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            rows = list(reader)
            headers = {str(header).strip() for header in reader.fieldnames or []}
        elif ext == "xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet.max_row > MAX_PLAN_VISITS + 1:
                raise TourPlanValidationError(
                    f"Upload no more than {MAX_PLAN_VISITS} visit rows at a time."
                )
            values = sheet.iter_rows(values_only=True)
            header_values = next(values, ())
            headers = {str(value).strip() for value in header_values if value is not None}
            headers_list = [str(value).strip() if value is not None else "" for value in header_values]
            rows = [
                {headers_list[index]: value for index, value in enumerate(row) if index < len(headers_list)}
                for row in values
                if any(value is not None for value in row)
            ]
    except TourPlanValidationError:
        raise
    except (UnicodeDecodeError, OSError, ValueError, BadZipFile, InvalidFileException, EOFError) as exc:
        raise TourPlanValidationError("Could not read the uploaded CSV/Excel file.") from exc
    finally:
        if workbook:
            workbook.close()
    missing = sorted(IMPORT_HEADERS - headers)
    if missing:
        raise TourPlanValidationError("Missing required file columns: " + ", ".join(missing))
    if not rows:
        raise TourPlanValidationError("The uploaded file contains no visit rows.")
    if len(rows) > MAX_PLAN_VISITS:
        raise TourPlanValidationError(f"Upload no more than {MAX_PLAN_VISITS} visit rows at a time.")
    return rows, []


def _validate_import_rows(
    db: Session,
    raw_rows: list[dict],
    territory_id: int,
    period_start: date,
    period_end: date,
) -> tuple[list[dict], list[dict]]:
    valid_rows: list[dict] = []
    errors: list[dict] = []
    seen: set[tuple[date, int]] = set()
    seen_sequences: set[tuple[date, int]] = set()
    for row_number, raw in enumerate(raw_rows, start=2):
        row = {str(key).strip(): value for key, value in raw.items()}
        try:
            parsed_date = _parse_date(str(row.get("planned_date", ""))[:10], "planned_date")
            doctor_id = parse_integer(row.get("doctor_id"))
            visit_type_id = parse_integer(row.get("visit_type_id"))
            raw_sequence = row.get("sequence")
            sequence = parse_integer(
                (row_number - 1) if raw_sequence in (None, "") else raw_sequence
            )
            if (parsed_date, doctor_id) in seen:
                raise TourPlanValidationError("Duplicate doctor visit on this date.")
            if (parsed_date, sequence) in seen_sequences:
                raise TourPlanValidationError("Sequence numbers must be unique for each date.")
            cleaned = validate_plan_visits(
                db,
                territory_id,
                period_start,
                period_end,
                [{
                    "planned_date": parsed_date.isoformat(),
                    "doctor_id": doctor_id,
                    "visit_type_id": visit_type_id,
                    "sequence": sequence,
                    "planned_products": row.get("planned_products"),
                    "planned_indication": row.get("planned_indication"),
                    "notes": row.get("notes"),
                }],
            )[0]
            seen.add((parsed_date, doctor_id))
            seen_sequences.add((parsed_date, sequence))
            valid_rows.append(
                {
                    **{key: (value.isoformat() if isinstance(value, date) else value) for key, value in cleaned.items()},
                    "row_number": row_number,
                }
            )
        except (TourPlanValidationError, TypeError, ValueError) as exc:
            errors.append({"row_number": row_number, "message": str(exc)})
    return valid_rows, errors


@router.get("", response_class=HTMLResponse)
async def tour_plan_list(
    request: Request,
    view: str = "list",
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    marketer_id: Optional[int] = None,
    territory_id: Optional[int] = None,
    plan_status: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = visible_plans_query(db, current_user)
    try:
        start = date.fromisoformat(start_date) if start_date else None
        end = date.fromisoformat(end_date) if end_date else None
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Filter dates must use YYYY-MM-DD.") from exc
    if current_user.role != "marketer" and marketer_id:
        query = query.filter(TourPlan.marketer_id == marketer_id)
    if territory_id:
        query = query.filter(TourPlan.territory_id == territory_id)
    if start:
        query = query.filter(TourPlan.period_end >= start)
    if end:
        query = query.filter(TourPlan.period_start <= end)
    if plan_status in PLAN_STATUSES:
        query = query.filter(TourPlan.status == plan_status)
    plans = query.order_by(TourPlan.period_start.desc(), TourPlan.id.desc()).all()
    for plan in plans:
        plan.actual_summary = _plan_counts(plan)
    template_query = db.query(TourPlanTemplate).options(joinedload(TourPlanTemplate.territory))
    if current_user.role == "marketer":
        template_query = template_query.filter(TourPlanTemplate.territory_id == current_user.territory_id)
    saved_templates = template_query.order_by(TourPlanTemplate.name).all()

    return templates.TemplateResponse(
        request=request,
        name="tour_plan/list.html",
        context={
            "user": current_user,
            "plans": plans,
            "view": "calendar" if view == "calendar" else "list",
            "start_date": start_date or "",
            "end_date": end_date or "",
            "selected_marketer_id": marketer_id,
            "selected_territory_id": territory_id,
            "selected_status": plan_status or "",
            "marketers": db.query(User).filter(User.role == "marketer", User.is_active.is_(True)).order_by(User.full_name).all() if current_user.role != "marketer" else [],
            "territories": db.query(Territory).order_by(Territory.name).all(),
            "can_create": current_user.role == "admin" or (current_user.role == "marketer" and _can_self_service(db)),
            "self_service": _can_self_service(db),
            "saved_templates": saved_templates,
        },
    )


@router.post("/settings")
async def update_tour_plan_settings(
    marketer_self_service: bool = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not _is_admin(current_user):
        raise HTTPException(status_code=403, detail="Only Admin can change Tour Plan self-service settings.")
    setting = get_tour_plan_setting(db)
    setting.marketer_self_service = marketer_self_service
    setting.updated_by_id = current_user.id
    db.commit()
    return RedirectResponse("/tour-plan", status_code=status.HTTP_303_SEE_OTHER)


def _plan_counts(plan: TourPlan) -> dict:
    rows = [plan_visit_actuals(visit) for visit in plan.visits]
    return {
        "planned": len(rows),
        "completed": sum(item["status"] == "Completed" for item in rows),
        "pending": sum(item["status"] == "Pending DCR" for item in rows),
        "not_visited": sum(item["status"] == "Not Visited" for item in rows),
        "percentage": round(100 * sum(item["status"] == "Completed" for item in rows) / len(rows)) if rows else 0,
    }


@router.get("/new", response_class=HTMLResponse)
async def new_tour_plan(
    request: Request,
    copy_from_id: Optional[int] = None,
    template_id: Optional[int] = None,
    template: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_can_create(db, current_user)
    copy_from = _get_plan_or_404(db, copy_from_id, current_user) if copy_from_id else None
    rows = None
    if template_id:
        saved_template = db.query(TourPlanTemplate).filter(TourPlanTemplate.id == template_id).first()
        if not saved_template:
            raise HTTPException(status_code=404, detail="Tour Plan template not found.")
        if current_user.role == "marketer" and current_user.territory_id != saved_template.territory_id:
            raise HTTPException(status_code=403, detail="Template is outside your assigned territory.")
        context = _plan_form_context(request, db, current_user)
        start = date.fromisoformat(context["period_start"])
        if saved_template.period_type == "monthly":
            start = date(start.year, start.month, 1)
            context["period_start"] = start.isoformat()
            context["period_end"] = date(
                start.year, start.month, monthrange(start.year, start.month)[1]
            ).isoformat()
        rows = [
            {
                **item,
                "planned_date": (start + timedelta(days=item["day_offset"])).isoformat(),
            }
            for item in json.loads(saved_template.visits_json)
        ]
        context = _plan_form_context(request, db, current_user, template_rows=rows)
        context["selected_territory_id"] = saved_template.territory_id
        context["period_type"] = saved_template.period_type
        context["title"] = saved_template.name
        context["period_start"] = start.isoformat()
        if saved_template.period_type == "monthly":
            context["period_end"] = date(
                start.year, start.month, monthrange(start.year, start.month)[1]
            ).isoformat()
        return templates.TemplateResponse(request=request, name="tour_plan/form.html", context=context)
    context = _plan_form_context(request, db, current_user, copy_from=copy_from, template_rows=rows)
    context["weekly_template"] = template == "weekly"
    if template == "weekly":
        start = date.fromisoformat(context["period_start"])
        context["rows"] = [
            {
                "planned_date": (start + timedelta(days=day)).isoformat(),
                "sequence": 1,
            }
            for day in range(5)
        ]
    return templates.TemplateResponse(request=request, name="tour_plan/form.html", context=context)


@router.post("", response_class=HTMLResponse)
async def create_tour_plan(
    request: Request,
    title: str = Form(..., max_length=150),
    period_type: str = Form(...),
    period_start: str = Form(...),
    period_end: str = Form(...),
    territory_id: int = Form(...),
    marketer_id: Optional[int] = Form(None),
    visits_json: str = Form(...),
    copy_from_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_can_create(db, current_user)
    assigned_marketer_id = marketer_id if _is_admin(current_user) else current_user.id
    if not assigned_marketer_id:
        raise HTTPException(status_code=400, detail="Select a marketer.")
    form_values = _submitted_form_values(
        title, period_type, period_start, period_end,
        territory_id, assigned_marketer_id, visits_json,
    )
    try:
        plan = create_plan(
            db,
            current_user,
            title,
            period_type,
            _parse_date(period_start, "Start date"),
            _parse_date(period_end, "End date"),
            assigned_marketer_id,
            territory_id,
            _parse_visits(visits_json),
            action="copied" if copy_from_id else "created",
        )
    except TourPlanValidationError as exc:
        context = _plan_form_context(
            request, db, current_user, error=str(exc), form_values=form_values
        )
        return templates.TemplateResponse(
            request=request, name="tour_plan/form.html", context=context, status_code=400
        )
    except IntegrityError as exc:
        db.rollback()
        context = _plan_form_context(
            request, db, current_user,
            error="Duplicate planned visit detected. Remove duplicate doctor/date rows and try again.",
            form_values=form_values,
        )
        return templates.TemplateResponse(
            request=request, name="tour_plan/form.html", context=context, status_code=400
        )
    return RedirectResponse(f"/tour-plan/{plan.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/import", response_class=HTMLResponse)
async def import_tour_plan_page(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_can_create(db, current_user)
    context = _plan_form_context(request, db, current_user)
    return templates.TemplateResponse(request=request, name="tour_plan/import.html", context=context)


@router.get("/import/template.csv")
async def download_import_template(current_user: User = Depends(get_current_user)):
    content = (
        "planned_date,sequence,doctor_id,visit_type_id,planned_products,planned_indication,notes\n"
        "2026-10-20,1,1,1,StemPro-X,Orthopedic / Osteoarthritis,Discuss clinical evidence\n"
    )
    return Response(
        content=content,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="tour-plan-template.csv"'},
    )


@router.post("/import/preview", response_class=HTMLResponse)
async def preview_tour_plan_import(
    request: Request,
    title: str = Form(..., max_length=150),
    period_type: str = Form(...),
    period_start: str = Form(...),
    period_end: str = Form(...),
    territory_id: int = Form(...),
    marketer_id: Optional[int] = Form(None),
    upload: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_can_create(db, current_user)
    assigned_marketer_id = marketer_id if _is_admin(current_user) else current_user.id
    if not assigned_marketer_id:
        raise HTTPException(status_code=400, detail="Select a marketer.")
    try:
        start = _parse_date(period_start, "Start date")
        end = _parse_date(period_end, "End date")
        if period_type not in {"weekly", "monthly"} or end < start:
            raise TourPlanValidationError("Plan period is invalid.")
        validate_plan_period(period_type, start, end)
        file_content = await upload.read(MAX_UPLOAD_BYTES + 1)
        raw_rows, _ = _decode_import_upload(upload.filename or "", file_content)
        valid_rows, errors = _validate_import_rows(db, raw_rows, territory_id, start, end)
        if not valid_rows:
            raise TourPlanValidationError("No valid visits to import. Correct the listed row errors and upload again.")
    except TourPlanValidationError as exc:
        context = _plan_form_context(request, db, current_user, error=str(exc))
        return templates.TemplateResponse(
            request=request, name="tour_plan/import.html", context=context, status_code=400
        )

    batch_id = str(uuid.uuid4())
    import_batch = TourPlanImport(
        id=batch_id,
        uploaded_by_id=current_user.id,
        metadata_json=json.dumps({
            "title": title.strip(),
            "period_type": period_type,
            "period_start": start.isoformat(),
            "period_end": end.isoformat(),
            "territory_id": territory_id,
            "marketer_id": assigned_marketer_id,
        }),
        rows_json=json.dumps(valid_rows),
        errors_json=json.dumps(errors),
    )
    db.add(import_batch)
    db.commit()
    log_context = {
        **_plan_form_context(request, db, current_user),
        "batch_id": batch_id,
        "title": title,
        "valid_rows": valid_rows,
        "errors": errors,
    }
    return templates.TemplateResponse(
        request=request, name="tour_plan/import_preview.html", context=log_context
    )


@router.post("/import/{batch_id}/confirm")
async def confirm_tour_plan_import(
    batch_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_can_create(db, current_user)
    batch = db.query(TourPlanImport).filter(
        TourPlanImport.id == batch_id,
        TourPlanImport.uploaded_by_id == current_user.id,
        TourPlanImport.created_at >= datetime.now(timezone.utc) - IMPORT_TTL,
    ).first()
    if not batch:
        raise HTTPException(status_code=404, detail="Import preview expired. Upload the file again.")
    metadata = json.loads(batch.metadata_json)
    if current_user.role == "marketer" and metadata["marketer_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Import is not assigned to your account.")
    visits = json.loads(batch.rows_json)
    try:
        plan = create_plan(
            db,
            current_user,
            metadata["title"],
            metadata["period_type"],
            date.fromisoformat(metadata["period_start"]),
            date.fromisoformat(metadata["period_end"]),
            metadata["marketer_id"],
            metadata["territory_id"],
            visits,
            action="imported",
        )
        db.delete(batch)
        db.commit()
    except (TourPlanValidationError, IntegrityError) as exc:
        db.rollback()
        logger.warning("Tour Plan import confirmation failed (%s).", type(exc).__name__)
        raise HTTPException(status_code=400, detail="Imported plan failed validation; re-upload and review the rows.") from exc
    return RedirectResponse(f"/tour-plan/{plan.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/report", response_class=HTMLResponse)
async def tour_plan_report(
    request: Request,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    marketer_id: Optional[int] = None,
    territory_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        start = date.fromisoformat(start_date) if start_date else None
        end = date.fromisoformat(end_date) if end_date else None
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Report dates must use YYYY-MM-DD.") from exc
    rows = get_plan_report_rows(db, current_user, marketer_id, territory_id, start, end)
    planned_count = len(rows)
    completed_count = sum(row["status"] == "Completed" for row in rows)
    not_visited_count = sum(row["status"] == "Not Visited" for row in rows)
    pending_count = sum(row["status"] == "Pending DCR" for row in rows)

    unplanned_query = db.query(FactDCR).filter(
        ~FactDCR.id.in_(
            db.query(TourPlanVisit.actual_dcr_id).filter(TourPlanVisit.actual_dcr_id.is_not(None))
        )
    )
    if current_user.role == "marketer":
        unplanned_query = unplanned_query.filter(FactDCR.user_id == current_user.id)
    elif marketer_id:
        unplanned_query = unplanned_query.filter(FactDCR.user_id == marketer_id)
    if territory_id:
        unplanned_query = unplanned_query.filter(FactDCR.territory_id == territory_id)
    if start:
        unplanned_query = unplanned_query.filter(FactDCR.report_date >= start)
    if end:
        unplanned_query = unplanned_query.filter(FactDCR.report_date <= end)
    unplanned_dcrs = unplanned_query.options(
        joinedload(FactDCR.user),
        joinedload(FactDCR.territory),
        joinedload(FactDCR.doctor),
        joinedload(FactDCR.visit_type),
    ).order_by(FactDCR.report_date.desc()).all()

    return templates.TemplateResponse(
        request=request,
        name="tour_plan/report.html",
        context={
            "request": request,
            "user": current_user,
            "rows": rows,
            "unplanned_dcrs": unplanned_dcrs,
            "planned_count": planned_count,
            "completed_count": completed_count,
            "not_visited_count": not_visited_count,
            "pending_count": pending_count,
            "completion_percentage": round(completed_count * 100 / planned_count) if planned_count else 0,
            "marketers": db.query(User).filter(User.role == "marketer", User.is_active.is_(True)).order_by(User.full_name).all() if current_user.role != "marketer" else [],
            "territories": db.query(Territory).order_by(Territory.name).all(),
            "selected_marketer_id": marketer_id,
            "selected_territory_id": territory_id,
            "start_date": start_date or "",
            "end_date": end_date or "",
        },
    )


@router.get("/{plan_id}", response_class=HTMLResponse)
async def tour_plan_detail(
    request: Request,
    plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    plan = _get_plan_or_404(db, plan_id, current_user)
    rows = [{"visit": visit, **plan_visit_actuals(visit)} for visit in plan.visits]
    return templates.TemplateResponse(
        request=request,
        name="tour_plan/detail.html",
        context={
            "request": request,
            "user": current_user,
            "plan": plan,
            "rows": rows,
            "counts": _plan_counts(plan),
            "audits": audit_trail(db, plan.id),
            "can_edit": current_user.role == "admin" or (
                current_user.role == "marketer"
                and plan.marketer_id == current_user.id
                and _can_self_service(db)
                and plan.status in {"Draft", "Returned"}
            ),
            "is_admin": _is_admin(current_user),
            "self_service": _can_self_service(db),
        },
    )


@router.get("/{plan_id}/edit", response_class=HTMLResponse)
async def edit_tour_plan_page(
    request: Request,
    plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    plan = _get_plan_or_404(db, plan_id, current_user)
    _require_plan_edit(db, plan, current_user)
    return templates.TemplateResponse(
        request=request,
        name="tour_plan/form.html",
        context=_plan_form_context(request, db, current_user, plan=plan),
    )


@router.post("/{plan_id}/edit", response_class=HTMLResponse)
async def update_tour_plan(
    request: Request,
    plan_id: int,
    title: str = Form(..., max_length=150),
    period_type: str = Form(...),
    period_start: str = Form(...),
    period_end: str = Form(...),
    territory_id: int = Form(...),
    marketer_id: Optional[int] = Form(None),
    visits_json: str = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    plan = _get_plan_or_404(db, plan_id, current_user)
    _require_plan_edit(db, plan, current_user)
    assigned_marketer_id = marketer_id if _is_admin(current_user) else current_user.id
    form_values = _submitted_form_values(
        title, period_type, period_start, period_end,
        territory_id, assigned_marketer_id or 0, visits_json,
    )
    try:
        updated = update_plan(
            db,
            plan,
            current_user,
            title,
            period_type,
            _parse_date(period_start, "Start date"),
            _parse_date(period_end, "End date"),
            assigned_marketer_id or 0,
            territory_id,
            _parse_visits(visits_json),
        )
    except TourPlanValidationError as exc:
        context = _plan_form_context(
            request, db, current_user, plan=plan, error=str(exc), form_values=form_values
        )
        return templates.TemplateResponse(
            request=request, name="tour_plan/form.html", context=context, status_code=400
        )
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail="Duplicate visit detected.") from exc
    return RedirectResponse(f"/tour-plan/{updated.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/{plan_id}/submit")
async def submit_plan_for_review(
    plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    plan = _get_plan_or_404(db, plan_id, current_user)
    if current_user.role != "marketer" or plan.marketer_id != current_user.id:
        raise HTTPException(status_code=403, detail="Only the assigned marketer can submit this plan.")
    if not _can_self_service(db):
        raise HTTPException(status_code=403, detail="Tour Plan self-service is disabled.")
    if plan.status not in {"Draft", "Returned"}:
        raise HTTPException(status_code=409, detail="Only draft or returned plans can be submitted.")
    if not plan.visits:
        raise HTTPException(status_code=400, detail="Add at least one planned visit before submission.")
    plan.status = "Submitted"
    log_plan_action(db, plan, current_user, "submitted")
    db.commit()
    return RedirectResponse(f"/tour-plan/{plan.id}", status_code=status.HTTP_303_SEE_OTHER)


async def _admin_plan_transition(
    request: Request,
    plan_id: int,
    transition: str,
    db: Session,
    current_user: User,
    note: str | None = None,
):
    if not _is_admin(current_user):
        raise HTTPException(status_code=403, detail="Only Admin can review, publish, or assign Tour Plans.")
    plan = _get_plan_or_404(db, plan_id, current_user)
    if transition == "approve":
        if plan.status != "Submitted":
            raise HTTPException(status_code=409, detail="Only submitted plans can be approved.")
        plan.status = "Approved"
        plan.approved_by_id = current_user.id
        plan.approved_at = datetime.now(timezone.utc)
    elif transition == "publish":
        if plan.status not in {"Draft", "Approved"}:
            raise HTTPException(
                status_code=409,
                detail="Only a draft or approved plan can be published.",
            )
        if not plan.visits:
            raise HTTPException(status_code=400, detail="Add at least one planned visit before publishing.")
        plan.status = "Published"
        plan.published_by_id = current_user.id
        plan.published_at = datetime.now(timezone.utc)
    elif transition == "return":
        if plan.status not in {"Submitted", "Approved"}:
            raise HTTPException(status_code=409, detail="Only submitted or approved plans can be returned.")
        if not note or not note.strip():
            raise HTTPException(status_code=400, detail="Add a reason when returning a plan.")
        plan.status = "Returned"
        plan.approved_by_id = None
        plan.approved_at = None
        plan.published_by_id = None
        plan.published_at = None
    else:
        raise HTTPException(status_code=400, detail="Unsupported plan action.")
    log_plan_action(db, plan, current_user, transition, note.strip() if note else None)
    db.commit()
    return RedirectResponse(f"/tour-plan/{plan.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/{plan_id}/approve")
async def approve_plan(
    request: Request, plan_id: int, db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await _admin_plan_transition(request, plan_id, "approve", db, current_user)


@router.post("/{plan_id}/publish")
async def publish_plan(
    request: Request, plan_id: int, db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await _admin_plan_transition(request, plan_id, "publish", db, current_user)


@router.post("/{plan_id}/return")
async def return_plan(
    request: Request,
    plan_id: int,
    note: str = Form(..., max_length=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await _admin_plan_transition(request, plan_id, "return", db, current_user, note)


@router.post("/{plan_id}/visits/{visit_id}/status")
async def update_planned_visit_status(
    plan_id: int,
    visit_id: int,
    visit_status: str = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if visit_status not in VISIT_STATUSES:
        raise HTTPException(status_code=400, detail="Invalid planned visit status.")
    plan = _get_plan_or_404(db, plan_id, current_user)
    if current_user.role != "admin" and (
        current_user.role != "marketer"
        or current_user.id != plan.marketer_id
        or not _can_self_service(db)
    ):
        raise HTTPException(status_code=403, detail="You cannot update visits on this plan.")
    visit = next((item for item in plan.visits if item.id == visit_id), None)
    if not visit:
        raise HTTPException(status_code=404, detail="Planned visit not found.")
    if visit_status == "Completed" and not visit.actual_dcr_id:
        raise HTTPException(status_code=409, detail="Only a submitted DCR can mark a planned visit completed.")
    if visit.actual_dcr_id and visit_status != "Completed":
        raise HTTPException(status_code=409, detail="A visit linked to a submitted DCR is completed.")
    visit.status = visit_status
    log_plan_action(db, plan, current_user, "visit_status_changed", f"Visit {visit.id}: {visit_status}")
    db.commit()
    return RedirectResponse(f"/tour-plan/{plan.id}", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/{plan_id}/templates")
async def create_tour_plan_template(
    plan_id: int,
    name: str = Form(..., max_length=120),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    plan = _get_plan_or_404(db, plan_id, current_user)
    if current_user.role != "admin" and (
        current_user.role != "marketer" or plan.marketer_id != current_user.id
        or not _can_self_service(db)
    ):
        raise HTTPException(status_code=403, detail="You cannot save another marketer's plan as a template.")
    if not name.strip():
        raise HTTPException(status_code=400, detail="Template name is required.")
    save_plan_template(db, plan, current_user, name)
    return RedirectResponse("/tour-plan", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/visits/options", response_class=HTMLResponse)
async def planned_visit_options(
    report_date: Optional[str] = None,
    doctor_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        requested_date = date.fromisoformat(report_date or "")
    except ValueError:
        return HTMLResponse("<option value=''>No matching planned visit</option>")
    query = db.query(TourPlanVisit).join(TourPlan).filter(
        TourPlan.marketer_id == current_user.id,
        TourPlan.status == "Published",
        TourPlanVisit.planned_date == requested_date,
        TourPlanVisit.actual_dcr_id.is_(None),
        TourPlanVisit.status.in_(["Planned", "Rescheduled"]),
    )
    if doctor_id:
        query = query.filter(TourPlanVisit.doctor_id == doctor_id)
    visits = query.options(
        joinedload(TourPlanVisit.doctor),
        joinedload(TourPlanVisit.plan),
        joinedload(TourPlanVisit.visit_type),
    ).order_by(TourPlanVisit.sequence).all()
    if not visits:
        return HTMLResponse("<option value=''>Unplanned call / automatic matching</option>")
    options = ["<option value=''>Unplanned / auto-match if unique</option>"]
    for visit in visits:
        label = (
            f"{visit.sequence}. {visit.doctor.name} - {visit.visit_type.name} "
            f"({visit.plan.title})"
        )
        options.append(f"<option value='{visit.id}'>{html.escape(label)}</option>")
    return HTMLResponse("\n".join(options))
